#!/usr/bin/env python3
"""Debug test for failing endpoints"""
import requests
import json
import tempfile
import os

BASE_URL = "http://localhost:5000"

def test_pptx():
    """Test PPTX conversion"""
    print("\n==== TESTING PPTX Conversion ====")
    
    # Create minimal PPTX file
    from pptx import Presentation
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    # Save to temp file
    with tempfile.NamedTemporaryFile(suffix='.pptx', delete=False) as f:
        prs.save(f.name)
        pptx_path = f.name
    
    try:
        # Get token first
        resp = requests.post(f"{BASE_URL}/api/auth/login", json={
            "username": "user1",
            "password": "pass123"
        })
        token = resp.json().get('token')
        
        # Test PPTX
        with open(pptx_path, 'rb') as f:
            files = {'file': f}
            headers = {'Authorization': f'Bearer {token}'}
            resp = requests.post(f"{BASE_URL}/api/conversions/pptx-to-pdf", 
                                files=files, headers=headers)
            
            print(f"Status: {resp.status_code}")
            try:
                print(f"Response: {json.dumps(resp.json(), indent=2)}")
            except:
                print(f"Response: {resp.text}")
    finally:
        os.unlink(pptx_path)

def test_merge():
    """Test merge endpoint"""
    print("\n==== TESTING Merge Endpoint ====")
    
    # Create two test PDFs
    from reportlab.pdfgen import canvas
    
    def create_pdf(filename):
        c = canvas.Canvas(filename)
        c.drawString(100, 750, "Test PDF")
        c.save()
    
    pdf1 = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
    pdf2 = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
    
    create_pdf(pdf1)
    create_pdf(pdf2)
    
    try:
        # Get token
        resp = requests.post(f"{BASE_URL}/api/auth/login", json={
            "username": "user1",
            "password": "pass123"
        })
        token = resp.json().get('token')
        
        # Test merge
        with open(pdf1, 'rb') as f1, open(pdf2, 'rb') as f2:
            files = [
                ('files', ('pdf1.pdf', f1, 'application/pdf')),
                ('files', ('pdf2.pdf', f2, 'application/pdf')),
            ]
            headers = {'Authorization': f'Bearer {token}'}
            resp = requests.post(f"{BASE_URL}/api/conversions/merge", 
                                files=files, headers=headers)
            
            print(f"Status: {resp.status_code}")
            try:
                print(f"Response: {json.dumps(resp.json(), indent=2)}")
            except:
                print(f"Response: {resp.text}")
    finally:
        os.unlink(pdf1)
        os.unlink(pdf2)

def test_batch():
    """Test batch endpoint"""
    print("\n==== TESTING Batch Endpoint ====")
    
    from reportlab.pdfgen import canvas
    
    def create_pdf(filename):
        c = canvas.Canvas(filename)
        c.drawString(100, 750, "Test PDF")
        c.save()
    
    pdf1 = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
    create_pdf(pdf1)
    
    try:
        # Get token
        resp = requests.post(f"{BASE_URL}/api/auth/login", json={
            "username": "user1",
            "password": "pass123"
        })
        token = resp.json().get('token')
        
        # Test batch
        with open(pdf1, 'rb') as f1:
            files = [('files', ('test.pdf', f1, 'application/pdf'))]
            data = {'tool_type': 'pdf-to-docx'}
            headers = {'Authorization': f'Bearer {token}'}
            resp = requests.post(f"{BASE_URL}/api/conversions/batch", 
                                files=files, data=data, headers=headers)
            
            print(f"Status: {resp.status_code}")
            try:
                print(f"Response: {json.dumps(resp.json(), indent=2)}")
            except:
                print(f"Response: {resp.text}")
    finally:
        os.unlink(pdf1)

def test_register():
    """Test register endpoint"""
    print("\n==== TESTING Register Endpoint ====")
    
    resp = requests.post(f"{BASE_URL}/api/auth/register", json={
        "username": "testuser123",
        "password": "test123pass",
        "email": "test@example.com"
    })
    
    print(f"Status: {resp.status_code}")
    try:
        print(f"Response: {json.dumps(resp.json(), indent=2)}")
    except:
        print(f"Response: {resp.text}")

if __name__ == '__main__':
    test_pptx()
    test_merge()
    test_batch()
    test_register()
