#!/usr/bin/env python3
"""Create sample test files for API testing"""
import os

print("Creating test files...\n")

# 1. Create PDF
try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    
    filename = "sample.pdf"
    if not os.path.exists(filename):
        c = canvas.Canvas(filename, pagesize=letter)
        c.drawString(100, 750, "Sample PDF Document")
        c.drawString(100, 700, "This is a test PDF for conversion testing.")
        c.drawString(100, 650, "Page 1")
        c.showPage()
        c.drawString(100, 750, "Page 2")
        c.save()
        print(f"✓ Created {filename}")
    else:
        print(f"✓ {filename} already exists")
except Exception as e:
    print(f"✗ Failed to create PDF: {e}")

# 2. Create DOCX
try:
    from docx import Document
    
    filename = "sample.docx"
    if not os.path.exists(filename):
        doc = Document()
        doc.add_heading("Sample Word Document", 0)
        doc.add_paragraph("This is a test document for conversion testing.")
        doc.add_paragraph("It contains multiple paragraphs.")
        doc.save(filename)
        print(f"✓ Created {filename}")
    else:
        print(f"✓ {filename} already exists")
except Exception as e:
    print(f"✗ Failed to create DOCX: {e}")

# 3. Create XLSX
try:
    from openpyxl import Workbook
    
    filename = "sample.xlsx"
    if not os.path.exists(filename):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Name"
        ws["B1"] = "Value"
        ws["C1"] = "Category"
        ws["A2"] = "Item 1"
        ws["B2"] = 100
        ws["C2"] = "A"
        ws["A3"] = "Item 2"
        ws["B3"] = 200
        ws["C3"] = "B"
        ws["A4"] = "Item 3"
        ws["B4"] = 300
        ws["C4"] = "A"
        wb.save(filename)
        print(f"✓ Created {filename}")
    else:
        print(f"✓ {filename} already exists")
except Exception as e:
    print(f"✗ Failed to create XLSX: {e}")

# 4. Create Image
try:
    from PIL import Image, ImageDraw
    
    filename = "sample.jpg"
    if not os.path.exists(filename):
        img = Image.new("RGB", (400, 300), color="lightblue")
        draw = ImageDraw.Draw(img)
        draw.text((150, 140), "Sample Image", fill="black")
        img.save(filename)
        print(f"✓ Created {filename}")
    else:
        print(f"✓ {filename} already exists")
except Exception as e:
    print(f"✗ Failed to create image: {e}")

# 5. Create CSV
filename = "sample.csv"
if not os.path.exists(filename):
    with open(filename, "w") as f:
        f.write("Name,Value,Category,Date\n")
        f.write("Product A,100,Electronics,2026-01-15\n")
        f.write("Product B,250,Electronics,2026-01-16\n")
        f.write("Product C,75,Software,2026-01-17\n")
        f.write("Product D,300,Hardware,2026-01-18\n")
    print(f"✓ Created {filename}")
else:
    print(f"✓ {filename} already exists")

# 6. Create PPTX
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    
    filename = "sample.pptx"
    if not os.path.exists(filename):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Sample Presentation"
        subtitle.text = "For Testing Conversions"
        
        # Add another slide
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        left = Inches(1)
        top = Inches(1)
        width = Inches(8)
        height = Inches(2)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.text = "This is slide 2 with sample content"
        
        prs.save(filename)
        print(f"✓ Created {filename}")
    else:
        print(f"✓ {filename} already exists")
except Exception as e:
    print(f"✗ Failed to create PPTX: {e}")

print("\n✅ Test files ready for API testing!")
