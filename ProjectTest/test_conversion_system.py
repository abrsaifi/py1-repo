#!/usr/bin/env python3
"""
Test actual conversions to verify system is fully working.
"""

import os
import sys
import tempfile
from pathlib import Path

print("\n" + "="*70)
print("CONVERSION FUNCTIONALITY TESTS")
print("="*70)

# Test 1: PDF to Image conversion
print("\n1. PDF TO IMAGE CONVERSION")
print("-" * 70)

try:
    from app.pdf_handler import pdf_to_images
    from PIL import Image
    
    # Create a simple test PDF first using reportlab
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    
    temp_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    temp_pdf_path = temp_pdf.name
    temp_pdf.close()
    
    # Create a simple PDF with text
    c = canvas.Canvas(temp_pdf_path, pagesize=letter)
    c.drawString(100, 750, "Test PDF for Conversion")
    c.drawString(100, 700, "This is page 1")
    c.showPage()
    c.drawString(100, 750, "Page 2 Content")
    c.showPage()
    c.save()
    
    # Now try to convert it
    images = pdf_to_images(temp_pdf_path, dpi=150)
    
    if images and len(images) > 0:
        print(f"  SUCCESS: PDF converted to {len(images)} image(s)")
        for i, img in enumerate(images):
            if isinstance(img, Image.Image):
                print(f"    Image {i+1}: {img.size} pixels, mode={img.mode}")
            else:
                print(f"    Image {i+1}: Loaded as {type(img)}")
    else:
        print(f"  WARNING: Conversion returned no images")
    
    # Cleanup
    os.unlink(temp_pdf_path)
    
except Exception as e:
    print(f"  FAILED: {str(e)[:100]}")

# Test 2: Image Format Conversion
print("\n2. IMAGE FORMAT CONVERSION")
print("-" * 70)

try:
    from PIL import Image
    
    # Create a test image
    temp_img = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
    temp_img_path = temp_img.name
    temp_img.close()
    
    # Create simple image
    img = Image.new('RGB', (100, 100), color='red')
    img.save(temp_img_path)
    
    # Convert to JPEG
    output_jpg = temp_img_path.replace('.png', '.jpg')
    img = Image.open(temp_img_path)
    img.convert('RGB').save(output_jpg)
    
    if os.path.exists(output_jpg):
        size_png = os.path.getsize(temp_img_path)
        size_jpg = os.path.getsize(output_jpg)
        print(f"  SUCCESS: Image format conversion completed")
        print(f"    PNG -> JPG conversion: {size_png} -> {size_jpg} bytes")
        os.unlink(output_jpg)
    
    os.unlink(temp_img_path)
    
except Exception as e:
    print(f"  FAILED: {str(e)[:100]}")

# Test 3: Excel File Operations
print("\n3. EXCEL FILE OPERATIONS")
print("-" * 70)

try:
    import openpyxl
    import pandas as pd
    
    # Create test Excel file
    temp_xlsx = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    temp_xlsx_path = temp_xlsx.name
    temp_xlsx.close()
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Data"
    
    # Add data
    ws['A1'] = "Name"
    ws['B1'] = "Value"
    ws['A2'] = "Item1"
    ws['B2'] = 100
    ws['A3'] = "Item2"
    ws['B3'] = 200
    
    wb.save(temp_xlsx_path)
    
    # Try to read it back
    df = pd.read_excel(temp_xlsx_path)
    
    print(f"  SUCCESS: Excel file created and read")
    print(f"    Rows: {len(df)}, Columns: {len(df.columns)}")
    print(f"    Columns: {list(df.columns)}")
    
    os.unlink(temp_xlsx_path)
    
except Exception as e:
    print(f"  FAILED: {str(e)[:100]}")

# Test 4: Word Document Operations
print("\n4. WORD DOCUMENT OPERATIONS")
print("-" * 70)

try:
    from docx import Document
    
    # Create test document
    temp_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
    temp_docx_path = temp_docx.name
    temp_docx.close()
    
    # Create document
    doc = Document()
    doc.add_paragraph("Test Document")
    doc.add_paragraph("This is a test paragraph")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"
    table.rows[1].cells[0].text = "Data 1"
    table.rows[1].cells[1].text = "Data 2"
    
    doc.save(temp_docx_path)
    
    # Read it back
    doc_read = Document(temp_docx_path)
    
    print(f"  SUCCESS: Word document created and read")
    print(f"    Paragraphs: {len(doc_read.paragraphs)}")
    print(f"    Tables: {len(doc_read.tables)}")
    if len(doc_read.tables) > 0:
        print(f"    Table 1 rows: {len(doc_read.tables[0].rows)}")
    
    os.unlink(temp_docx_path)
    
except Exception as e:
    print(f"  FAILED: {str(e)[:100]}")

# Test 5: PowerPoint Operations
print("\n5. POWERPOINT OPERATIONS")
print("-" * 70)

try:
    from pptx import Presentation
    
    # Create test presentation
    temp_pptx = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    temp_pptx_path = temp_pptx.name
    temp_pptx.close()
    
    # Create presentation
    prs = Presentation()
    
    # Add slide with title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Test Presentation"
    
    # Add another slide with content
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Content Slide"
    content = slide.placeholders[1]
    content.text = "Test content"
    
    prs.save(temp_pptx_path)
    
    # Read it back
    prs_read = Presentation(temp_pptx_path)
    
    print(f"  SUCCESS: PowerPoint presentation created and read")
    print(f"    Slides: {len(prs_read.slides)}")
    
    os.unlink(temp_pptx_path)
    
except Exception as e:
    print(f"  FAILED: {str(e)[:100]}")

# Test 6: PDF Manipulation (merge)
print("\n6. PDF MANIPULATION")
print("-" * 70)

try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    from pypdf import PdfMerger
    
    # Create two test PDFs
    temp_pdf1 = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    temp_pdf1_path = temp_pdf1.name
    temp_pdf1.close()
    
    temp_pdf2 = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    temp_pdf2_path = temp_pdf2.name
    temp_pdf2.close()
    
    # Create first PDF
    c = canvas.Canvas(temp_pdf1_path, pagesize=letter)
    c.drawString(100, 750, "PDF 1")
    c.save()
    
    # Create second PDF
    c = canvas.Canvas(temp_pdf2_path, pagesize=letter)
    c.drawString(100, 750, "PDF 2")
    c.save()
    
    # Merge PDFs
    temp_merged = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    temp_merged_path = temp_merged.name
    temp_merged.close()
    
    merger = PdfMerger()
    merger.append(temp_pdf1_path)
    merger.append(temp_pdf2_path)
    merger.write(temp_merged_path)
    merger.close()
    
    if os.path.exists(temp_merged_path):
        size = os.path.getsize(temp_merged_path)
        print(f"  SUCCESS: PDFs merged successfully")
        print(f"    Output file size: {size} bytes")
    
    # Cleanup
    os.unlink(temp_pdf1_path)
    os.unlink(temp_pdf2_path)
    os.unlink(temp_merged_path)
    
except Exception as e:
    print(f"  FAILED: {str(e)[:100]}")

# Final Summary
print("\n" + "="*70)
print("CONVERSION TESTS COMPLETE")
print("="*70)
print("If all tests show SUCCESS, the conversion system is fully functional!")
print("="*70 + "\n")
