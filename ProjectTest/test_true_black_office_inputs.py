import io

import pytest
from docx import Document
from openpyxl import Workbook

from server import app


@pytest.fixture
def client():
    app.config['TESTING'] = True
    with app.test_client() as c:
        yield c


def _make_docx_bytes() -> io.BytesIO:
    doc = Document()
    doc.add_heading('True Black Test', level=1)
    doc.add_paragraph('This is a test DOCX document with some text.')
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def _make_xlsx_bytes() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['A1'] = 'True Black Test'
    ws['A2'] = 'Row 2'
    ws['B2'] = 123
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def test_true_black_accepts_docx(client):
    data = {
        'dpi': '200',
        'threshold': '250',
        'contrast': '3.0',
        'sharpness': '2.5',
        'files': (_make_docx_bytes(), 'test.docx'),
    }

    resp = client.post('/convert', data=data, content_type='multipart/form-data')
    assert resp.status_code == 200
    assert resp.data[:4] == b'%PDF'


def test_true_black_accepts_xlsx(client):
    data = {
        'dpi': '200',
        'threshold': '250',
        'contrast': '3.0',
        'sharpness': '2.5',
        'files': (_make_xlsx_bytes(), 'test.xlsx'),
    }

    resp = client.post('/convert', data=data, content_type='multipart/form-data')
    assert resp.status_code == 200
    assert resp.data[:4] == b'%PDF'
