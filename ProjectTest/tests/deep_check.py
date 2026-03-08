import os
import tempfile
import shutil
import traceback

from pathlib import Path

import sys, os
# ensure project root is on sys.path so `import server` works when running from tests/
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import server


def make_image(path, text="Hello PDF", size=(400, 200)):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new('RGB', size, color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    try:
        f = ImageFont.load_default()
    except Exception:
        f = None
    d.text((10, 10), text, fill=(0, 0, 0), font=f)
    img.save(path)


def make_docx(path, text="Sample docx text"):
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


def make_xlsx(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Col1", "Col2", "Col3"])
    ws.append([1, 2, 3])
    ws.append(["a", "b", "c"])
    wb.save(path)


def make_pdf_via_reportlab(path, text="Sample PDF"):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    c = canvas.Canvas(path, pagesize=letter)
    c.drawString(100, 700, text)
    c.showPage()
    c.save()


def run_tests():
    # Set to True to run OCR on PDF pages (can be slow); keep False for quicker checks
    RUN_PDF_OCR = False
    temp_dir = tempfile.mkdtemp()
    print('Temp dir:', temp_dir)
    results = []

    try:
        # Create sample files
        img_path = os.path.join(temp_dir, 'sample.png')
        make_image(img_path, "Test OCR 123")

        docx_path = os.path.join(temp_dir, 'sample.docx')
        make_docx(docx_path, "This is a test document.")

        xlsx_path = os.path.join(temp_dir, 'sample.xlsx')
        make_xlsx(xlsx_path)

        pdf_path = os.path.join(temp_dir, 'sample.pdf')
        make_pdf_via_reportlab(pdf_path, "PDF for testing")

        # image_to_pdf
        print('-> image_to_pdf')
        out_image_pdf = os.path.join(temp_dir, 'image_converted.pdf')
        try:
            ok = server.image_to_pdf(img_path, out_image_pdf)
            results.append(('image_to_pdf', ok, out_image_pdf))
        except Exception as e:
            results.append(('image_to_pdf', False, str(e)))

        # docx_to_pdf
        print('-> docx_to_pdf')
        out_docx_pdf = os.path.join(temp_dir, 'docx_converted.pdf')
        try:
            ok = server.docx_to_pdf(docx_path, out_docx_pdf)
            results.append(('docx_to_pdf', ok, out_docx_pdf))
        except Exception as e:
            results.append(('docx_to_pdf', False, str(e)))

        # excel_to_pdf
        print('-> excel_to_pdf')
        out_xlsx_pdf = os.path.join(temp_dir, 'xlsx_converted.pdf')
        try:
            ok = server.excel_to_pdf(xlsx_path, out_xlsx_pdf)
            results.append(('excel_to_pdf', ok, out_xlsx_pdf))
        except Exception as e:
            results.append(('excel_to_pdf', False, str(e)))

        # pdf_to_true_bw
        print('-> pdf_to_true_bw')
        out_bw = os.path.join(temp_dir, 'bw.pdf')
        try:
            ok = server.pdf_to_true_bw(pdf_path, out_bw, dpi=150, threshold=200)
            results.append(('pdf_to_true_bw', ok, out_bw))
        except Exception as e:
            results.append(('pdf_to_true_bw', False, str(e)))

        # encrypt & decrypt
        print('-> encrypt & decrypt')
        enc = os.path.join(temp_dir, 'enc.pdf')
        dec = os.path.join(temp_dir, 'dec.pdf')
        try:
            ok_enc = server.encrypt_pdf(pdf_path, enc, 'pass123')
            ok_dec, msg = server.decrypt_pdf(enc, dec, 'pass123')
            results.append(('encrypt_pdf', ok_enc, enc))
            results.append(('decrypt_pdf', ok_dec and msg == '', dec))
        except Exception as e:
            results.append(('encrypt/decrypt', False, str(e)))

        # split & merge
        print('-> split & merge')
        try:
            splits = server.split_pdf(pdf_path, temp_dir, [(1,1)])
            results.append(('split_pdf', bool(splits), splits))
            if splits:
                merged = os.path.join(temp_dir, 'merged.pdf')
                ok = server.merge_pdf(splits, merged)
                results.append(('merge_pdf', ok, merged))
        except Exception as e:
            results.append(('split/merge', False, str(e)))

        # remove pages
        print('-> remove pages')
        try:
            rem = os.path.join(temp_dir, 'removed.pdf')
            ok = server.remove_pages_from_pdf(pdf_path, [1], rem)
            results.append(('remove_pages_from_pdf', ok, rem))
        except Exception as e:
            results.append(('remove_pages_from_pdf', False, str(e)))

        # ocr_extract_text (image)
        print('-> ocr_extract_text (image)')
        try:
            ocr_out = os.path.join(temp_dir, 'ocr.txt')
            ok = server.ocr_extract_text(img_path, ocr_out)
            results.append(('ocr_extract_text_image', ok, ocr_out))
        except Exception as e:
            results.append(('ocr_extract_text_image', False, str(e)))

        # ocr_extract_text (pdf)
        if RUN_PDF_OCR:
            print('-> ocr_extract_text (pdf)')
            try:
                ocr_pdf_out = os.path.join(temp_dir, 'ocr_pdf.txt')
                ok = server.ocr_extract_text(pdf_path, ocr_pdf_out)
                results.append(('ocr_extract_text_pdf', ok, ocr_pdf_out))
            except Exception as e:
                results.append(('ocr_extract_text_pdf', False, str(e)))
        else:
            results.append(('ocr_extract_text_pdf', 'SKIPPED', 'skipped by RUN_PDF_OCR'))

    finally:
        print('\nResults:')
        for name, ok, info in results:
            print(f"{name}: {'OK' if ok else 'FAIL'} -> {info}")
        # cleanup
        try:
            shutil.rmtree(temp_dir)
        except Exception:
            pass


if __name__ == '__main__':
    run_tests()
