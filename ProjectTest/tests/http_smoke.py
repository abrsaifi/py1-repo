import requests
import os
import tempfile
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import server

BASE = 'http://127.0.0.1:5000'

def make_sample_files(tmpdir):
    from PIL import Image, ImageDraw
    # image
    img = Image.new('RGB', (200,100), color=(255,255,255))
    d = ImageDraw.Draw(img)
    d.text((10,10), 'hello', fill=(0,0,0))
    img_path = os.path.join(tmpdir, 'sample.png')
    img.save(img_path)

    # docx
    from docx import Document
    doc = Document()
    doc.add_paragraph('sample doc')
    docx_path = os.path.join(tmpdir, 'sample.docx')
    doc.save(docx_path)

    # xlsx
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(['a','b','c'])
    xlsx_path = os.path.join(tmpdir, 'sample.xlsx')
    wb.save(xlsx_path)

    # pdf
    from reportlab.pdfgen import canvas
    pdf_path = os.path.join(tmpdir, 'sample.pdf')
    c = canvas.Canvas(pdf_path)
    c.drawString(100,750,'Sample PDF')
    c.save()

    return img_path, docx_path, xlsx_path, pdf_path


def test_convert_to_pdf(tmpdir):
    img, docx, xlsx, pdf = make_sample_files(tmpdir)
    files = [
        ('file', open(img,'rb')),
        ('file', open(docx,'rb')),
        ('file', open(xlsx,'rb')),
    ]
    r = requests.post(BASE + '/convert-to-pdf', files=files)
    print('/convert-to-pdf', r.status_code, r.headers.get('content-type'))


def test_convert_bw(tmpdir):
    _,_,_,pdf = make_sample_files(tmpdir)
    files = [('pdf_file', open(pdf,'rb'))]
    r = requests.post(BASE + '/convert', files=files, data={'dpi':'150','threshold':'200'})
    print('/convert', r.status_code, r.headers.get('content-type'))


def test_pdf_extract(tmpdir):
    _,_,_,pdf = make_sample_files(tmpdir)
    files = [('files', open(pdf,'rb'))]
    r = requests.post(BASE + '/pdf-extract', files=files, data={'format':'docx'})
    print('/pdf-extract', r.status_code, r.headers.get('content-type'))


def main():
    tmpdir = tempfile.mkdtemp()
    try:
        test_convert_to_pdf(tmpdir)
        test_convert_bw(tmpdir)
        test_pdf_extract(tmpdir)
    finally:
        pass

if __name__ == '__main__':
    main()
