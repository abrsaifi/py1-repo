import csv
import json
import os
import sqlite3
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _read_dataframe(input_path):
    suffix = Path(input_path).suffix.lower()
    if suffix == '.csv':
        return pd.read_csv(input_path)
    if suffix in {'.xlsx', '.xls'}:
        return pd.read_excel(input_path)
    if suffix == '.json':
        with open(input_path, 'r', encoding='utf-8-sig') as handle:
            payload = json.load(handle)
        if isinstance(payload, list):
            return pd.DataFrame(payload)
        if isinstance(payload, dict):
            return pd.json_normalize(payload)
    raise ValueError(f'unsupported_dataset_format:{suffix}')


def dataframe_to_workbook(dataframe, output_path, sheet_name='Sheet1'):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31] or 'Sheet1'
    worksheet.append(list(dataframe.columns))
    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append(list(row))
    workbook.save(output_path)
    workbook.close()
    return True


def csv_to_excel_file(input_path, output_path):
    dataframe = pd.read_csv(input_path)
    return dataframe_to_workbook(dataframe, output_path, sheet_name=Path(input_path).stem)


def formulas_to_values_file(input_path, output_path):
    source_workbook = load_workbook(input_path)
    values_workbook = load_workbook(input_path, data_only=True)

    try:
        for sheet_name in source_workbook.sheetnames:
            source_sheet = source_workbook[sheet_name]
            values_sheet = values_workbook[sheet_name]
            for row in source_sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith('='):
                        replacement = values_sheet[cell.coordinate].value
                        cell.value = replacement if replacement is not None else cell.value

        source_workbook.save(output_path)
        return True
    finally:
        source_workbook.close()
        values_workbook.close()


def clean_charts_workbook(input_path, output_path):
    workbook = load_workbook(input_path)
    try:
        for worksheet in workbook.worksheets:
            if hasattr(worksheet, '_charts'):
                worksheet._charts = []
        workbook.save(output_path)
        return True
    finally:
        workbook.close()


def _apply_missing_value_strategy(dataframe, strategy):
    if strategy == 'skip':
        return dataframe.dropna()

    numeric_columns = dataframe.select_dtypes(include=['number']).columns
    if strategy == 'zero':
        dataframe[numeric_columns] = dataframe[numeric_columns].fillna(0)
    elif strategy == 'median':
        dataframe[numeric_columns] = dataframe[numeric_columns].fillna(dataframe[numeric_columns].median(numeric_only=True))
    else:
        dataframe[numeric_columns] = dataframe[numeric_columns].fillna(dataframe[numeric_columns].mean(numeric_only=True))
    return dataframe


def normalize_dataset_file(input_path, output_path, method='min-max', handle_missing='mean', round_decimals=4, output_format='xlsx'):
    dataframe = _read_dataframe(input_path).copy()
    dataframe = _apply_missing_value_strategy(dataframe, handle_missing)

    numeric_columns = list(dataframe.select_dtypes(include=['number']).columns)
    for column in numeric_columns:
        series = dataframe[column].astype(float)
        if method == 'z-score':
            std_dev = series.std()
            dataframe[column] = 0 if std_dev in (0, None) or pd.isna(std_dev) else (series - series.mean()) / std_dev
        elif method == 'decimal':
            max_abs = series.abs().max()
            if max_abs and not pd.isna(max_abs):
                scale = len(str(int(max_abs)))
                dataframe[column] = series / (10 ** scale)
        else:
            min_value = series.min()
            max_value = series.max()
            if max_value == min_value:
                dataframe[column] = 0
            else:
                dataframe[column] = (series - min_value) / (max_value - min_value)

    dataframe = dataframe.round(int(round_decimals))
    if output_format == 'csv':
        dataframe.to_csv(output_path, index=False)
        return True
    return dataframe_to_workbook(dataframe, output_path, sheet_name='Normalized')


def dataset_validation_report(input_path):
    dataframe = _read_dataframe(input_path)
    total_cells = max(len(dataframe), 1) * max(len(dataframe.columns), 1)
    missing_values = int(dataframe.isnull().sum().sum())
    duplicate_rows = int(len(dataframe) - len(dataframe.drop_duplicates()))
    quality_score = max(0.0, 100.0 - ((missing_values / total_cells) * 100.0) - ((duplicate_rows / max(len(dataframe), 1)) * 100.0))
    return {
        'generated_at': datetime.now(UTC).isoformat(),
        'total_rows': int(len(dataframe)),
        'total_columns': int(len(dataframe.columns)),
        'columns': [str(column) for column in dataframe.columns],
        'missing_values': {str(column): int(count) for column, count in dataframe.isnull().sum().to_dict().items()},
        'data_types': {str(column): str(dtype) for column, dtype in dataframe.dtypes.to_dict().items()},
        'duplicate_rows': duplicate_rows,
        'quality_score': round(quality_score, 2),
    }


def write_json_file(payload, output_path):
    with open(output_path, 'w', encoding='utf-8') as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False)
        handle.write('\n')
    return True


def export_dataframe_pdf(dataframe, output_path, title):
    document = SimpleDocTemplate(output_path, pagesize=letter)
    styles = getSampleStyleSheet()
    rows = [list(dataframe.columns)] + dataframe.head(25).fillna('').astype(str).values.tolist()
    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4b99')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d0d7e2')),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f8fafc')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ]))
    story = [Paragraph(title, styles['Heading1']), Spacer(1, 12), table]
    document.build(story)
    return True


def pdf_export_file(input_path, output_path, title=None):
    dataframe = _read_dataframe(input_path)
    return export_dataframe_pdf(dataframe, output_path, title or f'Data Export: {Path(input_path).stem}')


def reporting_output_file(input_path, output_path, report_type='summary', output_format='pdf'):
    dataframe = _read_dataframe(input_path)
    numeric_summary = dataframe.describe(include='all').transpose().fillna('')
    if output_format == 'html':
        html = ['<html><body>', f'<h1>Report: {Path(input_path).stem}</h1>', f'<p>Type: {report_type}</p>', numeric_summary.to_html(), '</body></html>']
        with open(output_path, 'w', encoding='utf-8') as handle:
            handle.write('\n'.join(html))
        return True
    if output_format == 'txt':
        with open(output_path, 'w', encoding='utf-8') as handle:
            handle.write(f'Report: {Path(input_path).stem}\n')
            handle.write(f'Type: {report_type}\n\n')
            handle.write(numeric_summary.to_string())
            handle.write('\n')
        return True
    return export_dataframe_pdf(numeric_summary.reset_index(), output_path, f'{report_type.title()} Report: {Path(input_path).stem}')


def database_export_file(input_path, output_path, target_format='csv'):
    dataframe = _read_dataframe(input_path)
    if target_format == 'json':
        records = dataframe.fillna('').to_dict(orient='records')
        return write_json_file(records, output_path)
    if target_format == 'xlsx':
        return dataframe_to_workbook(dataframe, output_path, sheet_name='DatabaseExport')
    dataframe.to_csv(output_path, index=False)
    return True


def split_sheet_outputs(input_path, output_dir, output_format='separate-files'):
    suffix = Path(input_path).suffix.lower()
    outputs = []

    if suffix == '.csv':
        dataframe = pd.read_csv(input_path)
        stem = Path(input_path).stem
        if output_format == 'csv':
            output_path = os.path.join(output_dir, f'{stem}.csv')
            dataframe.to_csv(output_path, index=False)
        else:
            output_path = os.path.join(output_dir, f'{stem}.xlsx')
            dataframe_to_workbook(dataframe, output_path, sheet_name=stem)
        outputs.append((Path(output_path).name, output_path))
        return outputs

    workbook = load_workbook(input_path, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            sheet_name = worksheet.title[:31] or 'Sheet'
            rows = list(worksheet.values)
            if rows:
                header = rows[0]
                body = rows[1:]
                dataframe = pd.DataFrame(body, columns=header)
            else:
                dataframe = pd.DataFrame()

            safe_name = ''.join(ch if ch.isalnum() or ch in {'-', '_'} else '_' for ch in sheet_name).strip('_') or 'Sheet'
            if output_format == 'csv':
                output_path = os.path.join(output_dir, f'{safe_name}.csv')
                dataframe.to_csv(output_path, index=False)
            else:
                output_path = os.path.join(output_dir, f'{safe_name}.xlsx')
                dataframe_to_workbook(dataframe, output_path, sheet_name=sheet_name)
            outputs.append((Path(output_path).name, output_path))
        return outputs
    finally:
        workbook.close()


def read_history_records(db_path, limit=25):
    if not db_path or not os.path.exists(db_path):
        return []

    with sqlite3.connect(db_path) as connection:
        cursor = connection.cursor()
        cursor.execute(
            'SELECT timestamp, operation, files, status, message FROM history ORDER BY id DESC LIMIT ?',
            (int(limit),),
        )
        rows = cursor.fetchall()

    records = []
    for timestamp, operation, files_json, status, message in rows:
        try:
            files = json.loads(files_json) if files_json else []
        except Exception:
            files = []
        records.append({
            'timestamp': timestamp,
            'operation': operation,
            'files': files,
            'status': status,
            'message': message,
        })
    return records