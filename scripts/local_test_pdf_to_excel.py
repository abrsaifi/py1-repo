from __future__ import annotations

from pathlib import Path
import sys


def main() -> None:
    # Uses the sample PDF already in the repo.
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    pdf_path = repo_root / "Inv-566.pdf"
    out_path = repo_root / "Inv-566.extracted.xlsx"

    import server

    ok = server.pdf_to_excel(str(pdf_path), str(out_path))
    print("pdf_to_excel ok:", ok)
    print("input:", pdf_path)
    print("output:", out_path)

    if not ok or not out_path.exists():
        raise SystemExit(1)

    # Print a compact preview of the sheet values so we can see if it's still mashed.
    from openpyxl import load_workbook

    wb = load_workbook(out_path)
    ws = wb.active

    print("sheet size:", ws.max_row, ws.max_column)

    max_r = min(ws.max_row, 60)
    max_c = min(ws.max_column, 12)

    for r in range(1, max_r + 1):
        row_vals: list[str] = []
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            row_vals.append("" if v is None else str(v).replace("\n", " ")[:60])
        if any(x.strip() for x in row_vals):
            print(f"{r:02d}:", " | ".join(row_vals))


if __name__ == "__main__":
    main()
