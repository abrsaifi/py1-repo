#!/usr/bin/env python3
"""Quick test of conversion services"""
import os
import sys
import tempfile
from pathlib import Path

# Try importing the server
try:
    from server import execute_service_conversion, SERVICE_TOOLS
    print("✓ Server imported successfully")
except Exception as e:
    print(f"✗ Failed to import server: {e}")
    sys.exit(1)

# Test files directory
test_dir = tempfile.mkdtemp()
print(f"\nTest directory: {test_dir}")

# Test data
test_services = [
    ('To PDF', 'test.txt', 'test_out.pdf'),
    ('PDF to B&W', 'test.pdf', 'test_bw.pdf'),
    ('Excel to CSV', 'test.xlsx', 'test.csv'),
    ('Image Compress', 'test.jpg', 'test_compressed.jpg'),
]

print("\n" + "="*60)
print("CONVERSION SERVICE TEST")
print("="*60)

# Create test files
print("\n[1] Creating test files...")

# Create a simple text file
txt_file = os.path.join(test_dir, 'test.txt')
with open(txt_file, 'w') as f:
    f.write("Hello World!\nThis is a test file.")
print(f"  ✓ Created {txt_file}")

# Try to create a simple PDF (using reportlab if available)
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    pdf_file = os.path.join(test_dir, 'test.pdf')
    c = canvas.Canvas(pdf_file, pagesize=letter)
    c.drawString(100, 750, "Test PDF")
    c.save()
    print(f"  ✓ Created {pdf_file}")
except Exception as e:
    print(f"  ! Could not create PDF: {e}")

# Try to create a simple image
try:
    from PIL import Image
    img_file = os.path.join(test_dir, 'test.jpg')
    img = Image.new('RGB', (100, 100), color='red')
    img.save(img_file)
    print(f"  ✓ Created {img_file}")
except Exception as e:
    print(f"  ! Could not create image: {e}")

# Try to create a simple Excel file
try:
    from openpyxl import Workbook
    xlsx_file = os.path.join(test_dir, 'test.xlsx')
    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Test'
    ws['B1'] = 'Data'
    wb.save(xlsx_file)
    print(f"  ✓ Created {xlsx_file}")
except Exception as e:
    print(f"  ! Could not create Excel: {e}")

# Test conversions
print("\n[2] Testing conversions...")
test_results = []

# Test Text to PDF
if os.path.exists(txt_file):
    output_file = os.path.join(test_dir, 'test_out.pdf')
    try:
        result = execute_service_conversion('To PDF', txt_file, output_file)
        status = "✓" if result and os.path.exists(output_file) else "✗"
        print(f"  {status} Text to PDF: {result}")
        test_results.append(('Text to PDF', result))
    except Exception as e:
        print(f"  ✗ Text to PDF error: {e}")
        test_results.append(('Text to PDF', False))

# Test Image Compress
if os.path.exists(os.path.join(test_dir, 'test.jpg')):
    output_file = os.path.join(test_dir, 'test_compressed.jpg')
    try:
        result = execute_service_conversion('Image Compress', 
                                          os.path.join(test_dir, 'test.jpg'), 
                                          output_file,
                                          quality=50)
        status = "✓" if result and os.path.exists(output_file) else "✗"
        print(f"  {status} Image Compress: {result}")
        test_results.append(('Image Compress', result))
    except Exception as e:
        print(f"  ✗ Image Compress error: {e}")
        test_results.append(('Image Compress', False))

# Test Excel to CSV
if os.path.exists(os.path.join(test_dir, 'test.xlsx')):
    output_file = os.path.join(test_dir, 'test.csv')
    try:
        result = execute_service_conversion('Excel to CSV',
                                          os.path.join(test_dir, 'test.xlsx'),
                                          output_file)
        status = "✓" if result and os.path.exists(output_file) else "✗"
        print(f"  {status} Excel to CSV: {result}")
        test_results.append(('Excel to CSV', result))
    except Exception as e:
        print(f"  ✗ Excel to CSV error: {e}")
        test_results.append(('Excel to CSV', False))

# Summary
print("\n" + "="*60)
print("TEST SUMMARY")
print("="*60)
passed = sum(1 for _, result in test_results if result)
total = len(test_results)
print(f"\nPassed: {passed}/{total}")
for service, result in test_results:
    status = "✓" if result else "✗"
    print(f"  {status} {service}")

print("\nTest directory: " + test_dir)
print("\nTo clean up, delete: " + test_dir)
