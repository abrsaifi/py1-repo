#!/usr/bin/env python3
"""Test DOCX to PDF conversion with advanced parameters"""

import sys
sys.path.insert(0, '/c/Users/dell/OneDrive/Documents/py1')

from docx import Document
from docx.shared import Inches, Pt
import os

# Create a test DOCX file
docx_path = "test_document.docx"
doc = Document()

# Add title
title = doc.add_heading('Sample Report', 0)

# Add some content
doc.add_paragraph('This is a sample Word document created for testing.')
doc.add_paragraph('It contains multiple sections with different formatting.')

# Add a table
table = doc.add_table(rows=3, cols=3)
table.style = 'Table Grid'
header_cells = table.rows[0].cells
header_cells[0].text = 'Column 1'
header_cells[1].text = 'Column 2'
header_cells[2].text = 'Column 3'

for i in range(1, 3):
    row_cells = table.rows[i].cells
    row_cells[0].text = f'Data {i}a'
    row_cells[1].text = f'Data {i}b'
    row_cells[2].text = f'Data {i}c'

# Add more text
doc.add_paragraph('This document will be converted to PDF with custom formatting.')

# Save the document
doc.save(docx_path)
print(f"Created test DOCX: {docx_path}")

# Now test conversion via API
import requests

response = requests.post('http://localhost:5000/api/convert',
    files={'files[]': open(docx_path, 'rb')},
    data={
        'tool_name': 'To PDF',
        'orientation': 'landscape',
        'paper_size': 'A4',
        'margin_top': '15',
        'margin_bottom': '15',
        'margin_left': '10',
        'margin_right': '10',
        'scale_factor': '90',
        'page_numbers': 'true',
        'compression': 'high'
    }
)

result = response.json()
print(f"\nConversion Response:")
print(f"  Status: {response.status_code}")
print(f"  Success: {result.get('success')}")
if result.get('success'):
    print(f"  Output files: {result.get('files')}")
    print("✅ DOCX to PDF conversion successful!")
else:
    print(f"  Error: {result.get('error')}")
