"""
Sheet Management Feature Tests
Test the new sheet listing, multi-sheet conversion, and CSV combination features
"""

import os
import tempfile
import shutil
from pathlib import Path

VERBOSE = __name__ == '__main__'


def log(*args, **kwargs):
    if VERBOSE:
        print(*args, **kwargs)

# Test data for creating sample files
SAMPLE_DATA_SHEET1 = [
    ['Date', 'Product', 'Quantity', 'Price'],
    ['2024-01-01', 'Laptop', 5, 1200],
    ['2024-01-02', 'Mouse', 25, 35],
    ['2024-01-03', 'Keyboard', 15, 85],
]

SAMPLE_DATA_SHEET2 = [
    ['Month', 'Revenue', 'Expenses', 'Profit'],
    ['January', 50000, 30000, 20000],
    ['February', 55000, 32000, 23000],
    ['March', 60000, 34000, 26000],
]

SAMPLE_CSV_DATA = """Customer,Email,Phone,Status
John Doe,john@example.com,555-0001,Active
Jane Smith,jane@example.com,555-0002,Active
Bob Johnson,bob@example.com,555-0003,Inactive
"""


def create_test_excel_file(filepath, num_sheets=2):
    """Create a test Excel file with multiple sheets"""
    from openpyxl import Workbook
    
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    # Create Sheet 1
    ws1 = wb.create_sheet('Sales Q1')
    for row_idx, row_data in enumerate(SAMPLE_DATA_SHEET1, 1):
        for col_idx, value in enumerate(row_data, 1):
            ws1.cell(row=row_idx, column=col_idx, value=value)
    
    # Create Sheet 2 if requested
    if num_sheets >= 2:
        ws2 = wb.create_sheet('Financial Summary')
        for row_idx, row_data in enumerate(SAMPLE_DATA_SHEET2, 1):
            for col_idx, value in enumerate(row_data, 1):
                ws2.cell(row=row_idx, column=col_idx, value=value)
    
    # Create Sheet 3 if requested
    if num_sheets >= 3:
        ws3 = wb.create_sheet('Data Analysis')
        ws3['A1'] = 'Analysis Results'
        ws3['A2'] = 'Total Revenue: $165,000'
        ws3['A3'] = 'Total Expenses: $96,000'
        ws3['A4'] = 'Total Profit: $69,000'
    
    wb.save(filepath)
    wb.close()
    log(f"✓ Created test Excel file: {filepath}")


def create_test_csv_files(directory, num_files=3):
    """Create test CSV files in a directory"""
    csv_files = []
    
    months = ['January', 'February', 'March']
    for i, month in enumerate(months[:num_files]):
        filepath = os.path.join(directory, f'{month.lower()}_data.csv')
        
        with open(filepath, 'w', newline='') as f:
            f.write(f"Month,Sales,Calls,Conversions\n")
            f.write(f"{month},${5000 * (i+1)},50,{10 + i*2}\n")
            f.write(f"{month},${6000 * (i+1)},55,{12 + i*2}\n")
        
        csv_files.append(filepath)
        log(f"✓ Created test CSV file: {filepath}")
    
    return csv_files


def test_list_sheets():
    """Test getting sheet information"""
    log("\n" + "="*50)
    log("TEST 1: List Sheets")
    log("="*50)
    
    try:
        from services.document_conversion import get_sheet_info
        
        temp_dir = tempfile.mkdtemp()
        test_file = os.path.join(temp_dir, 'test.xlsx')
        
        create_test_excel_file(test_file, num_sheets=3)
        
        result = get_sheet_info(test_file)
        
        if result.get('success') or 'sheets' in result:
            log(f"✓ File: {result['file_name']}")
            log(f"✓ Type: {result['file_type']}")
            log(f"✓ Sheets found: {len(result['sheets'])}")
            
            for sheet in result['sheets']:
                log(f"  - Name: {sheet['name']}, Rows: {sheet['rows']}, Cols: {sheet['columns']}")
                if sheet.get('preview'):
                    log(f"    Preview: {sheet['preview'][0] if sheet['preview'] else 'N/A'}")
            
            log("✓ TEST PASSED: Sheet information retrieved successfully")
        else:
            log(f"✗ TEST FAILED: {result.get('error', 'Unknown error')}")
        assert result.get('success') or 'sheets' in result, result.get('error', 'Unknown error')
    
    except Exception as e:
        log(f"✗ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        raise AssertionError(f"Sheet info test failed: {e}") from e
    
    finally:
        try:
            shutil.rmtree(temp_dir)
        except:
            pass


def test_excel_to_pdf_sheets():
    """Test converting specific sheets to PDF"""
    log("\n" + "="*50)
    log("TEST 2: Excel to PDF (Specific Sheets)")
    log("="*50)
    
    try:
        from services.document_conversion import excel_to_pdf
        
        temp_dir = tempfile.mkdtemp()
        test_file = os.path.join(temp_dir, 'test.xlsx')
        output_pdf = os.path.join(temp_dir, 'output.pdf')
        
        create_test_excel_file(test_file, num_sheets=3)
        
        # Test 1: Convert single sheet
        result = excel_to_pdf(
            test_file, 
            output_pdf,
            sheets='Sales Q1',
            orientation='landscape',
            margin_top=20,
            margin_left=20
        )
        
        if result and os.path.exists(output_pdf):
            log(f"✓ Successfully converted 'Sales Q1' sheet")
            log(f"✓ PDF size: {os.path.getsize(output_pdf)} bytes")
            os.unlink(output_pdf)
        else:
            log("✗ Single sheet conversion failed")
            assert False, "Single sheet conversion failed"
        
        # Test 2: Convert multiple sheets merged
        result = excel_to_pdf(
            test_file,
            output_pdf,
            sheets=['Sales Q1', 'Financial Summary'],
            merge_sheets=True,
            orientation='portrait'
        )
        
        if result and os.path.exists(output_pdf):
            log(f"✓ Successfully merged 2 sheets into PDF")
            log(f"✓ PDF size: {os.path.getsize(output_pdf)} bytes")
            os.unlink(output_pdf)
        else:
            log("✗ Merge sheets conversion failed")
            assert False, "Merge sheets conversion failed"
        
        log("✓ TEST PASSED: Sheet conversion works correctly")
    
    except Exception as e:
        log(f"✗ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        raise AssertionError(f"Excel to PDF sheet test failed: {e}") from e
    
    finally:
        try:
            shutil.rmtree(temp_dir)
        except:
            pass


def test_combine_csvs():
    """Test combining multiple CSV files into Excel"""
    print("\n" + "="*50)
    print("TEST 3: Combine CSV Files")
    print("="*50)
    
    try:
        from services.document_conversion import combine_csvs_to_excel
        
        temp_dir = tempfile.mkdtemp()
        csv_files = create_test_csv_files(temp_dir, num_files=3)
        output_path = os.path.join(temp_dir, 'combined.xlsx')
        
        result = combine_csvs_to_excel(
            csv_files,
            output_path,
            sheet_names=['January Sales', 'February Sales', 'March Sales']
        )
        
        if result and os.path.exists(output_path):
            print(f"✓ Successfully combined {len(csv_files)} CSV files")
            print(f"✓ Excel file size: {os.path.getsize(output_path)} bytes")
            
            # Verify sheets were created
            from openpyxl import load_workbook
            wb = load_workbook(output_path)
            print(f"✓ Created {len(wb.sheetnames)} sheets: {wb.sheetnames}")
            wb.close()
            
            print("✓ TEST PASSED: CSV combining works correctly")
        else:
            print("✗ CSV combining failed")
        assert result and os.path.exists(output_path), "CSV combining failed"
    
    except Exception as e:
        print(f"✗ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        raise AssertionError(f"CSV combining test failed: {e}") from e
    
    finally:
        try:
            shutil.rmtree(temp_dir)
        except:
            pass


def test_csv_sheet_info():
    """Test getting sheet info from CSV"""
    print("\n" + "="*50)
    print("TEST 4: CSV Sheet Information")
    print("="*50)
    
    try:
        from services.document_conversion import get_sheet_info
        
        temp_dir = tempfile.mkdtemp()
        csv_file = os.path.join(temp_dir, 'test.csv')
        
        # Write test CSV
        with open(csv_file, 'w') as f:
            f.write(SAMPLE_CSV_DATA)
        
        result = get_sheet_info(csv_file)
        
        if result.get('success') or 'sheets' in result:
            sheet = result['sheets'][0]
            print(f"✓ File: {result['file_name']}")
            print(f"✓ Type: {result['file_type']}")
            print(f"✓ Rows: {sheet['rows']}, Columns: {sheet['columns']}")
            print(f"✓ Preview: {sheet['preview'][:2]}")
            
            print("✓ TEST PASSED: CSV sheet info retrieved successfully")
        else:
            print(f"✗ TEST FAILED: {result.get('error', 'Unknown error')}")
        assert result.get('success') or 'sheets' in result, result.get('error', 'Unknown error')
    
    except Exception as e:
        print(f"✗ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        raise AssertionError(f"CSV sheet info test failed: {e}") from e
    
    finally:
        try:
            shutil.rmtree(temp_dir)
        except:
            pass


def run_all_tests():
    """Run all sheet management tests"""
    print("\n" + "="*70)
    print("SHEET MANAGEMENT FEATURE TEST SUITE")
    print("="*70)
    
    tests = [
        ('List Sheets', test_list_sheets),
        ('Excel to PDF (Sheets)', test_excel_to_pdf_sheets),
        ('Combine CSVs', test_combine_csvs),
        ('CSV Sheet Info', test_csv_sheet_info),
    ]
    
    results = []
    for test_name, test_func in tests:
        try:
            test_func()
            results.append((test_name, True))
        except Exception as e:
            print(f"\n✗ CRITICAL ERROR in {test_name}: {e}")
            results.append((test_name, False))
    
    # Summary
    print("\n" + "="*70)
    print("TEST SUMMARY")
    print("="*70)
    
    total = len(results)
    passed = sum(1 for _, result in results if result)
    failed = total - passed
    
    for test_name, result in results:
        status = "✓ PASSED" if result else "✗ FAILED"
        print(f"{status}: {test_name}")
    
    print(f"\nTotal: {total} | Passed: {passed} | Failed: {failed}")
    
    if failed == 0:
        print("\n✓ ALL TESTS PASSED!")
        return True
    else:
        print(f"\n✗ {failed} TEST(S) FAILED")
        return False


if __name__ == '__main__':
    success = run_all_tests()
    exit(0 if success else 1)
