from io import BytesIO
import subprocess
import zipfile

import imageio_ffmpeg
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter

from app import create_app


def _build_app(tmp_path):
    return create_app({
        'TESTING': True,
        'SQLALCHEMY_DATABASE_URI': 'sqlite:///:memory:',
        'SECRET_KEY': 'upload-test-secret',
        'JWT_SECRET_KEY': 'upload-test-jwt-secret-key-0123456789',
        'ENABLE_BACKGROUND_TASKS': False,
        'UPLOAD_API_KEY': 'test-upload-key',
        'UPLOAD_CHUNKS_DIR': str(tmp_path / 'chunks'),
        'UPLOAD_MAX_FILE_SIZE': 2 * 1024 * 1024,
        'UPLOAD_CLEANUP_RETENTION': 1,
    })


def _make_png_bytes():
    image = Image.new('RGB', (16, 16), (255, 0, 0))
    buffer = BytesIO()
    image.save(buffer, format='PNG')
    return buffer.getvalue()


def _make_pdf_bytes(page_count=1):
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _make_text_bytes(content='hello from text tool\n'):
    return content.encode('utf-8')


def _make_json_bytes():
    return b'{"hello":"world","count":1}'


def _make_csv_bytes():
    return b'name,value\nalpha,10\nbeta,20\n'


def _build_media_fixture(tmp_path, extension):
    ffmpeg_path = imageio_ffmpeg.get_ffmpeg_exe()
    media_dir = tmp_path / 'media-fixtures'
    media_dir.mkdir(exist_ok=True)

    wav_path = media_dir / 'sample.wav'
    if not wav_path.exists():
        subprocess.run(
            [
                ffmpeg_path,
                '-y',
                '-f',
                'lavfi',
                '-i',
                'sine=frequency=440:duration=1',
                '-ac',
                '1',
                '-ar',
                '22050',
                str(wav_path),
            ],
            check=True,
            capture_output=True,
        )

    mp3_path = media_dir / 'sample.mp3'
    if not mp3_path.exists():
        subprocess.run(
            [
                ffmpeg_path,
                '-y',
                '-i',
                str(wav_path),
                '-codec:a',
                'libmp3lame',
                '-q:a',
                '4',
                str(mp3_path),
            ],
            check=True,
            capture_output=True,
        )

    mp4_path = media_dir / 'sample.mp4'
    if not mp4_path.exists():
        subprocess.run(
            [
                ffmpeg_path,
                '-y',
                '-f',
                'lavfi',
                '-i',
                'testsrc=size=160x120:rate=12:duration=1',
                '-f',
                'lavfi',
                '-i',
                'sine=frequency=660:duration=1',
                '-shortest',
                '-pix_fmt',
                'yuv420p',
                '-c:v',
                'libx264',
                '-c:a',
                'aac',
                str(mp4_path),
            ],
            check=True,
            capture_output=True,
        )

    webm_path = media_dir / 'sample.webm'
    if not webm_path.exists():
        subprocess.run(
            [
                ffmpeg_path,
                '-y',
                '-i',
                str(mp4_path),
                '-c:v',
                'libvpx-vp9',
                '-b:v',
                '0',
                '-crf',
                '36',
                '-c:a',
                'libopus',
                str(webm_path),
            ],
            check=True,
            capture_output=True,
        )

    return {
        'wav': wav_path,
        'mp3': mp3_path,
        'mp4': mp4_path,
        'webm': webm_path,
    }[extension].read_bytes()


def test_upload_chunk_status_and_convert_flow(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    image_bytes = _make_png_bytes()
    first_chunk = image_bytes[: len(image_bytes) // 2]
    second_chunk = image_bytes[len(image_bytes) // 2 :]

    response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'upload-1',
            'filename': 'sample.png',
            'index': '0',
            'total': '2',
            'chunk': (BytesIO(first_chunk), 'chunk-0.part'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert response.status_code == 200
    payload = response.get_json()
    assert payload['success'] is True
    assert payload['assembled'] is False

    status_response = client.get('/api/upload-status?upload_id=upload-1', headers=headers)
    assert status_response.status_code == 200
    status_payload = status_response.get_json()
    assert status_payload['chunks'] == [0]
    assert status_payload['total'] == 2

    response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'upload-1',
            'filename': 'sample.png',
            'index': '1',
            'total': '2',
            'chunk': (BytesIO(second_chunk), 'chunk-1.part'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert response.status_code == 200
    payload = response.get_json()
    assert payload['success'] is True
    assert payload['assembled'] is True
    assert payload['assembled_path'].endswith('sample.png')

    convert_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'uploads': [{'upload_id': 'upload-1', 'filename': 'sample.png'}],
            'target_format': 'jpg',
            'quality': 80,
        },
    )
    assert convert_response.status_code == 200
    disposition = convert_response.headers.get('Content-Disposition', '')
    assert 'attachment' in disposition
    assert 'sample_converted.jpg' in disposition
    assert convert_response.data[:2] == b'\xff\xd8'


def test_upload_routes_require_api_key(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()

    response = client.get('/api/upload-status?upload_id=missing')
    assert response.status_code == 401
    assert response.get_json()['error'] == 'unauthorized'


def test_upload_chunk_requires_api_key(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()

    response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'upload-no-auth',
            'filename': 'sample.png',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_png_bytes()), 'sample.png'),
        },
        content_type='multipart/form-data',
    )

    assert response.status_code == 401
    assert response.get_json()['error'] == 'unauthorized'


def test_convert_uploaded_requires_api_key(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()

    response = client.post(
        '/api/convert-uploaded',
        json={'uploads': [], 'target_format': 'png'},
    )

    assert response.status_code == 401
    assert response.get_json()['error'] == 'unauthorized'


def test_convert_uploaded_rejects_invalid_json_payload(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    response = client.post(
        '/api/convert-uploaded',
        data='{',
        headers=headers,
        content_type='application/json',
    )

    assert response.status_code == 400
    assert response.get_json()['error'] == 'invalid_json'


def test_convert_uploaded_uses_parameters_payload_for_target_and_quality(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    image_bytes = _make_png_bytes()

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'upload-params',
            'filename': 'sample.png',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(image_bytes), 'sample.png'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200
    assert upload_response.get_json()['assembled'] is True

    convert_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'uploads': [{'upload_id': 'upload-params', 'filename': 'sample.png'}],
            'parameters': {
                'format': 'webp',
                'quality': 70,
            },
        },
    )

    assert convert_response.status_code == 200
    disposition = convert_response.headers.get('Content-Disposition', '')
    assert 'sample_converted.webp' in disposition
    assert convert_response.data[:4] == b'RIFF'


def test_convert_uploaded_requires_target_for_generic_conversions(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'missing-target',
            'filename': 'sample.png',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_png_bytes()), 'sample.png'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'uploads': [{'upload_id': 'missing-target', 'filename': 'sample.png'}],
            'parameters': {},
        },
    )

    assert response.status_code == 400
    assert response.get_json()['error'] == 'missing_target'


def test_convert_uploaded_rejects_unsupported_target(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'unsupported-target',
            'filename': 'sample.png',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_png_bytes()), 'sample.png'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'uploads': [{'upload_id': 'unsupported-target', 'filename': 'sample.png'}],
            'target_format': 'exe',
            'parameters': {},
        },
    )

    assert response.status_code == 400
    assert response.get_json()['error'] == 'unsupported_target'


def test_convert_uploaded_returns_no_converted_files_for_missing_upload(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'uploads': [{'upload_id': 'missing-upload', 'filename': 'sample.png'}],
            'target_format': 'jpg',
            'parameters': {},
        },
    )

    assert response.status_code == 500
    assert response.get_json()['error'] == 'no_converted_files'


def test_special_tool_returns_missing_uploads_for_missing_upload(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'text-formatter',
            'uploads': [{'upload_id': 'missing-upload', 'filename': 'notes.txt'}],
            'target_format': 'txt',
            'parameters': {},
        },
    )

    assert response.status_code == 400
    assert response.get_json()['error'] == 'missing_uploads'


def test_admin_purge_uploads_removes_old_directories(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_dir = tmp_path / 'chunks' / 'old-upload'
    upload_dir.mkdir(parents=True)
    (upload_dir / 'meta.json').write_text('{}', encoding='utf-8')

    response = client.post('/api/admin/purge-uploads', headers=headers, json={'older_than': 0})
    assert response.status_code == 200
    payload = response.get_json()
    assert payload['success'] is True
    assert payload['removed'] >= 1
    assert not upload_dir.exists()


def test_convert_uploaded_merges_multiple_pdf_uploads(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    for index, name in enumerate(['first.pdf', 'second.pdf'], start=1):
        upload_response = client.post(
            '/api/upload-chunk',
            data={
                'upload_id': f'merge-{index}',
                'filename': name,
                'index': '0',
                'total': '1',
                'chunk': (BytesIO(_make_pdf_bytes()), name),
            },
            headers=headers,
            content_type='multipart/form-data',
        )
        assert upload_response.status_code == 200
        assert upload_response.get_json()['assembled'] is True

    merge_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'merge-pdf-smart',
            'uploads': [
                {'upload_id': 'merge-1', 'filename': 'first.pdf'},
                {'upload_id': 'merge-2', 'filename': 'second.pdf'},
            ],
            'parameters': {
                'page_size': 'auto',
            },
        },
    )

    assert merge_response.status_code == 200
    disposition = merge_response.headers.get('Content-Disposition', '')
    assert 'first_merged.pdf' in disposition
    assert len(PdfReader(BytesIO(merge_response.data)).pages) == 2


def test_convert_uploaded_splits_pdf_into_zip_archive(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'split-1',
            'filename': 'source.pdf',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_pdf_bytes(page_count=2)), 'source.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200
    assert upload_response.get_json()['assembled'] is True

    split_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'split-pdf',
            'uploads': [{'upload_id': 'split-1', 'filename': 'source.pdf'}],
            'parameters': {
                'split_mode': 'individual',
            },
        },
    )

    assert split_response.status_code == 200
    assert 'application/zip' in split_response.headers.get('Content-Type', '')
    archive = zipfile.ZipFile(BytesIO(split_response.data))
    assert sorted(archive.namelist()) == ['source_part_1.pdf', 'source_part_2.pdf']


def test_convert_uploaded_encrypts_pdf_with_password(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'encrypt-1',
            'filename': 'secure.pdf',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_pdf_bytes()), 'secure.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200
    assert upload_response.get_json()['assembled'] is True

    encrypt_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'encrypt-pdf-aes256',
            'uploads': [{'upload_id': 'encrypt-1', 'filename': 'secure.pdf'}],
            'parameters': {
                'password': 'top-secret',
            },
        },
    )

    assert encrypt_response.status_code == 200
    disposition = encrypt_response.headers.get('Content-Disposition', '')
    assert 'secure_encrypted.pdf' in disposition
    assert PdfReader(BytesIO(encrypt_response.data)).is_encrypted is True


def test_convert_uploaded_extracts_requested_pages(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'extract-1',
            'filename': 'extract-source.pdf',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_pdf_bytes(page_count=3)), 'extract-source.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    extract_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'extract-pdf',
            'uploads': [{'upload_id': 'extract-1', 'filename': 'extract-source.pdf'}],
            'parameters': {'pages': '2-3'},
        },
    )

    assert extract_response.status_code == 200
    assert len(PdfReader(BytesIO(extract_response.data)).pages) == 2


def test_convert_uploaded_removes_requested_pages(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'remove-1',
            'filename': 'remove-source.pdf',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_pdf_bytes(page_count=3)), 'remove-source.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    remove_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'remove-pages-range',
            'uploads': [{'upload_id': 'remove-1', 'filename': 'remove-source.pdf'}],
            'parameters': {'pages_to_remove': '2'},
        },
    )

    assert remove_response.status_code == 200
    assert len(PdfReader(BytesIO(remove_response.data)).pages) == 2


def test_convert_uploaded_watermarks_pdf_when_text_is_provided(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'watermark-1',
            'filename': 'watermark-source.pdf',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_pdf_bytes()), 'watermark-source.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    watermark_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'watermark-image-text',
            'uploads': [{'upload_id': 'watermark-1', 'filename': 'watermark-source.pdf'}],
            'parameters': {'watermark_text': 'CONFIDENTIAL', 'opacity': 0.3, 'font_size': 36},
        },
    )

    assert watermark_response.status_code == 200
    assert 'watermark-source_watermarked.pdf' in watermark_response.headers.get('Content-Disposition', '')


def test_convert_uploaded_redacts_pdf_when_keywords_are_provided(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'redact-1',
            'filename': 'redact-source.pdf',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_pdf_bytes()), 'redact-source.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    redact_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'redact-pdf-permanent',
            'uploads': [{'upload_id': 'redact-1', 'filename': 'redact-source.pdf'}],
            'parameters': {'keywords': 'secret,internal'},
        },
    )

    assert redact_response.status_code == 200
    assert 'redact-source_redacted.pdf' in redact_response.headers.get('Content-Disposition', '')


def test_convert_uploaded_decrypts_pdf_with_password(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt('top-secret')
    encrypted_buffer = BytesIO()
    writer.write(encrypted_buffer)
    encrypted_buffer.seek(0)

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'decrypt-1',
            'filename': 'locked.pdf',
            'index': '0',
            'total': '1',
            'chunk': (encrypted_buffer, 'locked.pdf'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    decrypt_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'decrypt-pdf-remove',
            'uploads': [{'upload_id': 'decrypt-1', 'filename': 'locked.pdf'}],
            'parameters': {'password': 'top-secret'},
        },
    )

    assert decrypt_response.status_code == 200
    assert PdfReader(BytesIO(decrypt_response.data)).is_encrypted is False


def test_convert_uploaded_supports_helper_backed_special_tools(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    for upload_id, filename, content in [
        ('bw-1', 'source.pdf', _make_pdf_bytes()),
        ('meta-1', 'meta.pdf', _make_pdf_bytes()),
        ('text-1', 'notes.txt', _make_text_bytes()),
        ('json-1', 'payload.json', _make_json_bytes()),
    ]:
        response = client.post(
            '/api/upload-chunk',
            data={
                'upload_id': upload_id,
                'filename': filename,
                'index': '0',
                'total': '1',
                'chunk': (BytesIO(content), filename),
            },
            headers=headers,
            content_type='multipart/form-data',
        )
        assert response.status_code == 200

    bw_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'pdf-to-bw-pro',
            'uploads': [{'upload_id': 'bw-1', 'filename': 'source.pdf'}],
            'target_format': 'pdf',
            'parameters': {'dpi': 150, 'threshold': 180},
        },
    )
    assert bw_response.status_code == 200
    assert 'source_bw.pdf' in bw_response.headers.get('Content-Disposition', '')

    metadata_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'metadata-clean',
            'uploads': [{'upload_id': 'meta-1', 'filename': 'meta.pdf'}],
            'target_format': 'pdf',
            'parameters': {},
        },
    )
    assert metadata_response.status_code == 200
    assert 'meta_metadata_cleaned.pdf' in metadata_response.headers.get('Content-Disposition', '')

    text_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'text-to-pdf',
            'uploads': [{'upload_id': 'text-1', 'filename': 'notes.txt'}],
            'target_format': 'pdf',
            'parameters': {},
        },
    )
    assert text_response.status_code == 200
    assert 'notes_converted.pdf' in text_response.headers.get('Content-Disposition', '')

    json_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'json-formatter',
            'uploads': [{'upload_id': 'json-1', 'filename': 'payload.json'}],
            'target_format': 'json',
            'parameters': {},
        },
    )
    assert json_response.status_code == 200
    assert b'"count": 1' in json_response.data


def test_convert_uploaded_supports_media_transcoding_targets(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_specs = [
        ('media-mp3', 'sample.mp3', _build_media_fixture(tmp_path, 'mp3')),
        ('media-mp4', 'sample.mp4', _build_media_fixture(tmp_path, 'mp4')),
        ('media-webm', 'sample.webm', _build_media_fixture(tmp_path, 'webm')),
    ]

    for upload_id, filename, content in upload_specs:
        response = client.post(
            '/api/upload-chunk',
            data={
                'upload_id': upload_id,
                'filename': filename,
                'index': '0',
                'total': '1',
                'chunk': (BytesIO(content), filename),
            },
            headers=headers,
            content_type='multipart/form-data',
        )
        assert response.status_code == 200
        assert response.get_json()['assembled'] is True

    wav_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'mp3-to-wav',
            'uploads': [{'upload_id': 'media-mp3', 'filename': 'sample.mp3'}],
            'target_format': 'wav',
        },
    )
    assert wav_response.status_code == 200
    assert wav_response.data[:4] == b'RIFF'

    webm_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'mp4-to-webm',
            'uploads': [{'upload_id': 'media-mp4', 'filename': 'sample.mp4'}],
            'target_format': 'webm',
        },
    )
    assert webm_response.status_code == 200
    assert webm_response.data[:4] == b'\x1aE\xdf\xa3'

    mp4_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'mp4-converter',
            'uploads': [{'upload_id': 'media-webm', 'filename': 'sample.webm'}],
            'target_format': 'mp4',
        },
    )
    assert mp4_response.status_code == 200
    assert b'ftyp' in mp4_response.data[:64]


def test_convert_uploaded_supports_csv_to_excel_and_sheet_tools(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    csv_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'csv-1',
            'filename': 'dataset.csv',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_csv_bytes()), 'dataset.csv'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert csv_response.status_code == 200

    xlsx_upload = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'sheet-1',
            'filename': 'sheet-source.xlsx',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_build_media_fixture(tmp_path, 'wav')[:0] + _build_workbook_fixture()), 'sheet-source.xlsx'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert xlsx_upload.status_code == 200

    csv_to_excel_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'csv-to-excel',
            'uploads': [{'upload_id': 'csv-1', 'filename': 'dataset.csv'}],
            'target_format': 'xlsx',
        },
    )
    assert csv_to_excel_response.status_code == 200
    workbook = load_workbook(BytesIO(csv_to_excel_response.data))
    assert workbook.active['A2'].value == 'alpha'
    workbook.close()

    split_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'split-sheets',
            'uploads': [{'upload_id': 'sheet-1', 'filename': 'sheet-source.xlsx'}],
            'target_format': 'zip',
            'parameters': {'output_format': 'separate-files'},
        },
    )
    assert split_response.status_code == 200
    archive = zipfile.ZipFile(BytesIO(split_response.data))
    assert sorted(archive.namelist()) == ['Details.xlsx', 'Summary.xlsx']


def _build_workbook_fixture():
    workbook = PdfWriter()
    del workbook
    from openpyxl import Workbook

    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append(['name', 'value'])
    summary.append(['alpha', 10])
    details = book.create_sheet('Details')
    details.append(['region', 'amount'])
    details.append(['north', 5])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def test_convert_uploaded_supports_reporting_validator_and_history_tools(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'report-1',
            'filename': 'metrics.csv',
            'index': '0',
            'total': '1',
            'chunk': (BytesIO(_make_csv_bytes()), 'metrics.csv'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    validator_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'data-validator',
            'uploads': [{'upload_id': 'report-1', 'filename': 'metrics.csv'}],
            'target_format': 'json',
        },
    )
    assert validator_response.status_code == 200
    assert b'"quality_score"' in validator_response.data

    reporting_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'reporting',
            'uploads': [{'upload_id': 'report-1', 'filename': 'metrics.csv'}],
            'target_format': 'pdf',
            'parameters': {'format': 'pdf', 'reportType': 'summary'},
        },
    )
    assert reporting_response.status_code == 200
    assert reporting_response.data[:4] == b'%PDF'

    history_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'convert-history',
            'uploads': [],
            'target_format': 'json',
        },
    )
    assert history_response.status_code == 200
    assert b'"records"' in history_response.data


def test_convert_uploaded_can_create_and_extract_zip_archives(tmp_path):
    app = _build_app(tmp_path)
    client = app.test_client()
    headers = {'X-API-Key': 'test-upload-key'}

    upload_specs = [
        ('zip-src-1', 'first.txt', _make_text_bytes('first\n')),
        ('zip-src-2', 'second.txt', _make_text_bytes('second\n')),
    ]
    for upload_id, filename, content in upload_specs:
        response = client.post(
            '/api/upload-chunk',
            data={
                'upload_id': upload_id,
                'filename': filename,
                'index': '0',
                'total': '1',
                'chunk': (BytesIO(content), filename),
            },
            headers=headers,
            content_type='multipart/form-data',
        )
        assert response.status_code == 200

    create_zip_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'create-zip',
            'uploads': [
                {'upload_id': 'zip-src-1', 'filename': 'first.txt'},
                {'upload_id': 'zip-src-2', 'filename': 'second.txt'},
            ],
            'target_format': 'zip',
            'parameters': {},
        },
    )
    assert create_zip_response.status_code == 200
    zip_reader = zipfile.ZipFile(BytesIO(create_zip_response.data))
    assert sorted(zip_reader.namelist()) == ['first.txt', 'second.txt']

    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('hello.txt', 'hello zip')
        archive.writestr('nested/world.txt', 'world zip')
    zip_buffer.seek(0)

    upload_response = client.post(
        '/api/upload-chunk',
        data={
            'upload_id': 'zip-extract-1',
            'filename': 'archive.zip',
            'index': '0',
            'total': '1',
            'chunk': (zip_buffer, 'archive.zip'),
        },
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 200

    extract_zip_response = client.post(
        '/api/convert-uploaded',
        headers=headers,
        json={
            'tool_slug': 'zip-extractor',
            'uploads': [{'upload_id': 'zip-extract-1', 'filename': 'archive.zip'}],
            'target_format': 'zip',
            'parameters': {},
        },
    )
    assert extract_zip_response.status_code == 200
    extracted_archive = zipfile.ZipFile(BytesIO(extract_zip_response.data))
    assert sorted(extracted_archive.namelist()) == ['hello.txt', 'world.txt']