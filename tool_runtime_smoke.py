import argparse
import json
import subprocess
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

import imageio_ffmpeg
import requests
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter


PRIORITY_TOOL_ORDER = {
    'split-pdf': 10,
    'split-pdf-batch': 11,
    'merge-pdf': 12,
    'merge-pdf-smart': 13,
}

ALIASES_COVERED_BY_CANONICAL = {
    'split-pdf-batch': 'split-pdf',
}


def build_temp_samples(root: Path):
    temp_dir = Path(tempfile.mkdtemp(prefix='docpro-tool-smoke-'))
    (temp_dir / 'sample.txt').write_text('hello from docpro smoke test\n', encoding='utf-8')
    (temp_dir / 'sample.html').write_text('<html><body><h1>DocPro</h1><p>smoke test</p></body></html>', encoding='utf-8')
    (temp_dir / 'sample.json').write_text('{"hello": "world", "n": 1}\n', encoding='utf-8')
    (temp_dir / 'sample.csv').write_text('name,value\nalpha,10\nbeta,20\n', encoding='utf-8')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Summary'
    worksheet.append(['name', 'value', 'formula'])
    worksheet.append(['alpha', 10, '=B2*2'])
    worksheet.append(['beta', 20, '=B3*2'])
    workbook.create_sheet('Details')
    workbook['Details'].append(['region', 'amount'])
    workbook['Details'].append(['north', 5])
    workbook['Details'].append(['south', 7])
    workbook.save(temp_dir / 'sample.xlsx')
    workbook.close()

    document = DocxDocument()
    document.add_heading('DocPro smoke sample', level=1)
    document.add_paragraph('This file is used by the runtime smoke test.')
    document.save(temp_dir / 'sample.docx')

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = 'DocPro Smoke Test'
    slide.placeholders[1].text = 'PPTX fixture for live runtime coverage.'
    presentation.save(temp_dir / 'sample.pptx')

    with zipfile.ZipFile(temp_dir / 'sample.zip', 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('hello.txt', 'hello zip')
        zf.writestr('nested/world.txt', 'world zip')
    ffmpeg_path = imageio_ffmpeg.get_ffmpeg_exe()
    subprocess.run(
        [
            ffmpeg_path,
            '-y',
            '-f',
            'lavfi',
            '-i',
            'sine=frequency=440:duration=1',
            '-ac',
            '1',
            '-ar',
            '22050',
            str(temp_dir / 'sample.wav'),
        ],
        check=True,
        capture_output=True,
    )
    subprocess.run(
        [
            ffmpeg_path,
            '-y',
            '-i',
            str(temp_dir / 'sample.wav'),
            '-codec:a',
            'libmp3lame',
            '-q:a',
            '4',
            str(temp_dir / 'sample.mp3'),
        ],
        check=True,
        capture_output=True,
    )
    subprocess.run(
        [
            ffmpeg_path,
            '-y',
            '-f',
            'lavfi',
            '-i',
            'testsrc=size=160x120:rate=12:duration=1',
            '-f',
            'lavfi',
            '-i',
            'sine=frequency=660:duration=1',
            '-shortest',
            '-pix_fmt',
            'yuv420p',
            '-c:v',
            'libx264',
            '-c:a',
            'aac',
            str(temp_dir / 'sample.mp4'),
        ],
        check=True,
        capture_output=True,
    )
    subprocess.run(
        [
            ffmpeg_path,
            '-y',
            '-i',
            str(temp_dir / 'sample.mp4'),
            '-c:v',
            'libvpx-vp9',
            '-b:v',
            '0',
            '-crf',
            '36',
            '-c:a',
            'libopus',
            str(temp_dir / 'sample.webm'),
        ],
        check=True,
        capture_output=True,
    )
    encrypted_pdf_path = temp_dir / 'encrypted.pdf'
    encrypted_writer = PdfWriter()
    encrypted_writer.add_blank_page(width=200, height=200)
    encrypted_writer.encrypt('top-secret')
    with encrypted_pdf_path.open('wb') as encrypted_handle:
        encrypted_writer.write(encrypted_handle)

    return {
        'pdf': root / 'sample.pdf',
        'docx': temp_dir / 'sample.docx',
        'pptx': temp_dir / 'sample.pptx',
        'xlsx': temp_dir / 'sample.xlsx',
        'csv': temp_dir / 'sample.csv',
        'jpg': root / 'sample.jpg',
        'jpeg': root / 'sample.jpg',
        'png': root / 'test_image.png',
        'webp': root / 'ProjectTest' / 'downloaded_test.webp',
        'txt': temp_dir / 'sample.txt',
        'html': temp_dir / 'sample.html',
        'htm': temp_dir / 'sample.html',
        'encrypted_pdf': encrypted_pdf_path,
        'json': temp_dir / 'sample.json',
        'zip': temp_dir / 'sample.zip',
        'mp3': temp_dir / 'sample.mp3',
        'wav': temp_dir / 'sample.wav',
        'mp4': temp_dir / 'sample.mp4',
        'webm': temp_dir / 'sample.webm',
    }


def get_default_parameters(tool):
    defaults = {}
    for param in tool.get('params') or []:
        if 'default' in param:
            defaults[param['name']] = param['default']
    return defaults


def get_normalized_target(tool, params, selected_file):
    output_format = tool.get('output_format')
    if output_format and output_format not in {'auto', 'variable'}:
        return ''.join(ch for ch in str(output_format).lower() if ch.isalnum())

    direct_target = params.get('output_format') or params.get('format')
    if direct_target:
        return ''.join(ch for ch in str(direct_target).lower() if ch.isalnum())

    tool_target = str(tool.get('to_format') or tool.get('toFormat') or '').lower()
    if 'pdf' in tool_target:
        return 'pdf'
    if 'docx' in tool_target:
        return 'docx'
    if 'pptx' in tool_target or 'powerpoint' in tool_target:
        return 'pptx'
    if 'xlsx' in tool_target or 'excel' in tool_target:
        return 'xlsx'
    if 'csv' in tool_target:
        return 'csv'
    if 'zip' in tool_target:
        return 'zip'
    if 'png' in tool_target:
        return 'png'
    if 'jpg' in tool_target or 'jpeg' in tool_target:
        return 'jpg'
    if 'webm' in tool_target:
        return 'webm'
    if 'webp' in tool_target:
        return 'webp'
    if 'tiff' in tool_target or 'tif' in tool_target:
        return 'tiff'
    if 'wav' in tool_target:
        return 'wav'
    if 'mp4' in tool_target:
        return 'mp4'
    if 'html' in tool_target:
        return 'html'
    if 'json' in tool_target:
        return 'json'
    if 'text' in tool_target:
        return 'txt'
    if tool.get('slug') == 'extract-extended':
        return 'txt'
    if tool.get('slug') == 'zip-extractor':
        return 'zip'
    if tool.get('slug') == 'split-sheets':
        return 'zip'
    if tool.get('slug') in {'data-validator', 'convert-history'}:
        return 'json'
    if tool.get('slug') == 'audio-converter':
        return 'wav'
    if tool.get('slug') in {'formulas-to-values', 'clean-charts', 'normalize-data'}:
        return 'xlsx'
    if tool.get('category') == 'image' and selected_file is not None:
        source_extension = selected_file.suffix.lower().lstrip('.')
        if source_extension in {'jpg', 'jpeg', 'png', 'bmp', 'gif', 'tiff', 'tif', 'webp'}:
            return 'tiff' if source_extension == 'tif' else source_extension
    return ''


def get_accepted_formats(tool):
    supported = tool.get('supported_formats') or []
    if supported:
        return [str(fmt).lower().lstrip('.') for fmt in supported]

    source_text = str(tool.get('from_format') or tool.get('fromFormat') or '').lower()
    if 'multiple files' in source_text:
        return ['txt']
    if 'multiple' in source_text:
        return ['docx', 'txt', 'html', 'xlsx', 'csv', 'pptx']
    if 'jpg' in source_text or 'jpeg' in source_text:
        return ['jpg', 'jpeg']
    if 'png' in source_text:
        return ['png']
    if 'webp' in source_text:
        return ['webp']
    if 'image' in source_text:
        return ['jpg', 'jpeg', 'png', 'webp', 'bmp', 'gif', 'tiff', 'tif']
    if 'pdf' in source_text:
        return ['pdf']
    if 'docx' in source_text:
        return ['docx']
    if 'pptx' in source_text or 'powerpoint' in source_text:
        return ['pptx']
    if 'xlsx' in source_text or 'excel' in source_text or 'spreadsheet' in source_text:
        return ['xlsx', 'csv']
    if 'html' in source_text:
        return ['html', 'htm']
    if 'text' in source_text:
        return ['txt']
    if 'document' in source_text:
        return ['docx', 'pdf', 'txt', 'html', 'pptx']
    if 'json' in source_text:
        return ['json']
    if 'zip' in source_text or 'archive' in source_text:
        return ['zip']
    if 'workbook' in source_text:
        return ['xlsx', 'csv']
    if 'dataset' in source_text or 'operational data' in source_text:
        return ['xlsx', 'csv', 'json']
    if 'database export' in source_text:
        return ['csv', 'json', 'xlsx']
    if 'mixed files' in source_text:
        return ['docx', 'txt', 'html', 'xlsx', 'csv', 'pdf']
    if 'completed jobs' in source_text:
        return []
    if 'mp3' in source_text or 'audio' in source_text:
        return ['mp3', 'wav']
    if 'mp4' in source_text or 'video' in source_text or 'webm' in source_text:
        return ['mp4', 'webm']
    return []


def pick_uploads(tool, samples):
    slug = tool.get('slug')
    if slug in {'merge-pdf', 'merge-pdf-smart'}:
        return [samples['pdf'], samples['pdf']]
    if slug in {'decrypt-pdf', 'decrypt-pdf-remove'}:
        return [samples['encrypted_pdf']]
    if slug in {'to-pdf', 'to-pdf-advanced'}:
        return [samples['docx']]
    if slug == 'pptx-to-pdf':
        return [samples['pptx']]
    if slug == 'csv-to-excel':
        return [samples['csv']]
    if slug in {'split-sheets', 'formulas-to-values', 'clean-charts'}:
        return [samples['xlsx']]
    if slug in {'normalize-data', 'data-validator', 'pdf-export', 'reporting', 'database-export'}:
        return [samples['csv']]
    if slug == 'audio-converter':
        return [samples['mp3']]
    if slug == 'bulk-convert':
        return [samples['txt'], samples['html']]
    if slug == 'mp4-converter':
        return [samples['webm']]
    if slug == 'mp4-to-webm':
        return [samples['mp4']]
    if slug == 'mp3-to-wav':
        return [samples['mp3']]
    if str(tool.get('from_format') or '').lower().startswith('multiple files'):
        return [samples['txt'], samples['json']]

    for ext in get_accepted_formats(tool):
        candidate = samples.get(ext)
        if candidate and candidate.exists():
            return [candidate]
    return []


def upload_file(api_base, path, slug):
    upload_id = f'{slug}-{uuid.uuid4().hex}'
    with path.open('rb') as handle:
        files = {'chunk': (path.name, handle, 'application/octet-stream')}
        data = {'upload_id': upload_id, 'filename': path.name, 'index': '0', 'total': '1'}
        response = requests.post(f'{api_base}/api/upload-chunk', files=files, data=data, timeout=30)
    response.raise_for_status()
    payload = response.json()
    return {'upload_id': payload['upload_id'], 'filename': payload['filename']}


def build_parameters(tool):
    parameters = get_default_parameters(tool)
    slug = tool['slug']
    if slug in {'extract-pdf', 'extract-extended'}:
        parameters.update({'pages': '1'})
        return parameters
    if slug in {'remove-pages', 'remove-pages-range'}:
        parameters.update({'pages_to_remove': '1'})
        return parameters
    if slug in {'encrypt-pdf', 'encrypt-pdf-aes256'}:
        parameters.update({'password': 'top-secret'})
        return parameters
    if slug in {'decrypt-pdf', 'decrypt-pdf-remove'}:
        parameters.update({'password': 'top-secret'})
        return parameters
    if slug in {'watermark-pdf', 'watermark-image-text'}:
        parameters.update({'watermark_text': 'CONFIDENTIAL', 'opacity': 0.3, 'font_size': 36})
        return parameters
    if slug in {'redact-pdf', 'redact-pdf-permanent'}:
        parameters.update({'keywords': 'secret,internal'})
        return parameters
    return parameters


def main():
    parser = argparse.ArgumentParser(description='Run a live DocPro tool runtime smoke check.')
    parser.add_argument('--frontend', default='http://127.0.0.1:5173')
    parser.add_argument('--backend', default='http://127.0.0.1:5060')
    parser.add_argument('--transport', choices=('frontend', 'backend'), default='frontend')
    parser.add_argument('--workspace', default=str(Path(__file__).resolve().parent))
    args = parser.parse_args()

    workspace = Path(args.workspace)
    samples = build_temp_samples(workspace)
    api_base = args.frontend if args.transport == 'frontend' else args.backend

    catalog = requests.get(f'{args.backend}/api/tools', timeout=30)
    catalog.raise_for_status()
    tools = catalog.json()['tools']
    tools = sorted(
        enumerate(tools),
        key=lambda item: (PRIORITY_TOOL_ORDER.get(item[1].get('slug'), 1000), item[0]),
    )

    summary = {
        'transport': args.transport,
        'catalog_count': len(tools),
        'passed': [],
        'failed': [],
        'covered_by_alias': [],
        'skipped': [],
    }

    for _, tool in tools:
        canonical_slug = ALIASES_COVERED_BY_CANONICAL.get(tool['slug'])
        if canonical_slug and any(record['slug'] == canonical_slug for record in summary['passed']):
            summary['covered_by_alias'].append({'slug': tool['slug'], 'canonical': canonical_slug})
            continue

        uploads = pick_uploads(tool, samples)
        parameters = build_parameters(tool)
        target = get_normalized_target(tool, parameters, uploads[0] if uploads else None)
        if not uploads and tool['slug'] != 'convert-history':
            summary['skipped'].append({'slug': tool['slug'], 'reason': 'no_sample_for_source'})
            continue
        if not target:
            summary['skipped'].append({'slug': tool['slug'], 'reason': 'no_target_format'})
            continue

        try:
            upload_payloads = [upload_file(api_base, path, tool['slug']) for path in uploads]
            request_payload = {
                'tool_slug': tool['slug'],
                'uploads': upload_payloads,
                'target_format': target,
                'parameters': parameters,
            }
            response = requests.post(
                f'{api_base}/api/convert-uploaded',
                json=request_payload,
                timeout=60,
            )
            if response.status_code == 500:
                try:
                    error_payload = response.json().get('error', {})
                except Exception:
                    error_payload = {}
                message = error_payload.get('message', '') if isinstance(error_payload, dict) else ''
                if 'Stream has ended unexpectedly' in message:
                    time.sleep(0.25)
                    response = requests.post(
                        f'{api_base}/api/convert-uploaded',
                        json=request_payload,
                        timeout=60,
                    )
            record = {
                'slug': tool['slug'],
                'target': target,
                'status': response.status_code,
                'content_type': response.headers.get('content-type'),
            }
            if response.ok:
                summary['passed'].append(record)
            else:
                try:
                    record['error'] = response.json().get('error')
                except Exception:
                    record['error'] = response.text[:200]
                summary['failed'].append(record)
        except Exception as exc:
            summary['failed'].append({'slug': tool['slug'], 'target': target, 'error': repr(exc)})

        time.sleep(0.05)

    print(json.dumps(summary, ensure_ascii=True, indent=2))
    raise SystemExit(1 if summary['failed'] else 0)


if __name__ == '__main__':
    main()